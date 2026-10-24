"""Cierre integrado sin proveedor externo: dinero, conservación, documentos e interfaz."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader
from streamlit.testing.v1 import AppTest

from copiloto import catalog, claude_engine as C, ficha as F
from copiloto.export import archivos
from copiloto.quote import cotizar


class TestRegresiones(unittest.TestCase):
    def test_respuesta_parcial_conserva_pendientes_y_contradicciones(self):
        f = F.nueva()
        f["contradicciones"] = ["No se ha resuelto cuántos usuarios necesitan correo."]
        f["preguntas_pendientes"] = [{"pregunta": "¿Cuántos usuarios?", "motivo": "Licencias"}]
        with patch.object(C, "ask", return_value=({"hechos": [], "candidatos": []}, {})), patch.object(F, "guardar"):
            F.procesar_mensaje(f, "Estamos revisando el dato")
        self.assertTrue(f["contradicciones"])
        self.assertTrue(f["preguntas_pendientes"])

    def test_respuesta_incompleta_no_borra_la_ficha(self):
        f = F.nueva()
        antes = deepcopy(f)
        with patch.object(C, "ask", return_value=({}, {})), self.assertRaises(ValueError):
            F.procesar_mensaje(f, "Hola")
        self.assertEqual(f, antes)

    def test_fallo_guardar_conserva_fecha_y_contenido(self):
        with TemporaryDirectory() as ruta, patch.object(F, "FICHAS", Path(ruta)):
            f = F.nueva()
            f["actualizado"] = "fecha anterior"
            antes = deepcopy(f)
            with patch.object(Path, "replace", side_effect=OSError("ocupado")), self.assertRaises(OSError):
                F.guardar(f)
            self.assertEqual(f, antes)
            self.assertEqual(list(Path(ruta).iterdir()), [])

    def test_plazo_guardado_invalida_confirmacion(self):
        with TemporaryDirectory() as ruta, patch.object(F, "FICHAS", Path(ruta)):
            f = F.nueva()
            F.confirmar(f, [{"sku": "SIM-M365-02", "cantidad": 10}])
            F.cambiar_plazo(f, 24)
            self.assertEqual(F.cargar(f["case_id"])["plazo"], 24)
            self.assertEqual(F.estado_confirmacion(f), "desactualizada")

    def test_confirmacion_y_plazo_son_un_solo_guardado(self):
        f = F.nueva()
        with patch.object(F, "guardar") as guardar:
            F.confirmar(f, [{"sku": "SIM-M365-02", "cantidad": 12}], plazo=24)
        guardar.assert_called_once()
        self.assertEqual(f["confirmacion"]["seleccion"][0]["cantidad"], 12)
        self.assertEqual(f["plazo"], 24)
        self.assertEqual(f["revision"], 1)
        self.assertEqual(F.estado_confirmacion(f), "vigente")

    def test_redondeo_decimal_y_extremos_no_rompen_el_cierre(self):
        q = cotizar([{"sku": "SIM-BK-01", "cantidad": 1.2345}])
        self.assertEqual(q["totales"]["mensual"], 2283.83)
        q = cotizar([{"sku": "SIM-BK-01", "cantidad": .1711}])
        self.assertEqual(q["totales"]["mensual"], 316.54)
        q = cotizar([{"sku": "SIM-AWS-01"}], consumo={"AWS": 1e100})
        self.assertTrue(q["incompleto"])
        self.assertEqual(q["totales"]["mensual"], 0)

    def test_acumulacion_no_definida_bloquea_borrador(self):
        q = cotizar([{"sku": "SIM-M365-02", "cantidad": 100}], plazo=24)
        self.assertTrue(any("acumulación" in b for b in q["bloqueos"]))

    def test_prompt_incluye_todas_las_columnas_y_cero_es_un_numero(self):
        texto = catalog.texto_para_claude()
        for columna in catalog.load()["productos"][0]:
            self.assertIn(f'"{columna}"', texto)
        self.assertEqual(C._numeros(0), {"0.0"})


class TestExportaciones(unittest.TestCase):
    def setUp(self):
        self.f = F.nueva("EXPORT-PRUEBA", empresa="Empresa & <prueba> | Norte")
        self.f["resumen_necesidad"] = "Respaldar correo sin perder el contexto del cliente."
        self.f["candidatos"] = [{"sku": "SIM-M365-02", "cantidad": 10}]
        with patch.object(F, "guardar"):
            F.confirmar(self.f, self.f["candidatos"])

    def test_cuatro_archivos_editables_con_identidad_y_mismo_importe(self):
        q = F.cotizacion(self.f, True)
        salida = archivos(self.f, q)
        self.assertEqual(len(salida), 4)
        for nombre, contenido in salida.items():
            with self.subTest(nombre=nombre):
                if nombre.endswith(".pdf"):
                    texto = "\n".join(p.extract_text() for p in PdfReader(BytesIO(contenido)).pages)
                else:
                    doc = Document(BytesIO(contenido))
                    texto = "\n".join(p.text for p in doc.paragraphs)
                    self.assertGreater(len(doc.tables), 0)
                self.assertIn("EXPORT-PRUEBA", texto)
                self.assertIn("revisión 0", texto)
                self.assertIn("DATOS SIMULADOS", texto)
                if "propuesta" in nombre:
                    self.assertIn("2,450.00", texto)

    def test_incompleto_solo_permite_expediente_y_revision_obsoleta_se_rechaza(self):
        self.f["revision"] += 1
        q = F.cotizacion(self.f, True)
        self.assertEqual(len(archivos(self.f, q)), 2)
        self.assertTrue(all("expediente" in n for n in archivos(self.f, q)))
        self.f["revision"] += 1
        with self.assertRaises(ValueError):
            archivos(self.f, q)

    def test_consumo_sin_cantidad_se_exporta_y_texto_largo_se_pagina(self):
        self.f["hechos"] = [{"clave": "consumo_aws_mxn", "valor": 20000, "verificado": True}]
        self.f["transcript"] = [{"rol": "cliente", "texto": "La operación necesita continuidad. " * 250}]
        with patch.object(F, "guardar"):
            F.confirmar(self.f, [{"sku": "SIM-AWS-01"}])
        salida = archivos(self.f, F.cotizacion(self.f, True))
        pdf_interno = next(v for k, v in salida.items() if k.endswith("expediente.pdf"))
        self.assertGreater(len(PdfReader(BytesIO(pdf_interno)).pages), 1)


class TestInterfaz(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.storage = patch.object(F, "FICHAS", Path(self.tmp.name))
        self.storage.start()
        self.proveedor = patch.object(C, "ask", side_effect=RuntimeError("proveedor no disponible"))
        self.mock = self.proveedor.start()
        self.app = AppTest.from_file(str(Path(__file__).parents[1] / "app.py"), default_timeout=20).run()

    def tearDown(self):
        self.proveedor.stop()
        self.storage.stop()
        self.tmp.cleanup()

    def boton(self, label):
        return next(b for b in self.app.button if b.label == label)

    def vista(self, nombre):
        self.app.radio(key="navegacion").set_value(nombre).run()
        self.assertFalse(self.app.exception)

    def test_catalogo_45_y_cierre_mixto_con_cuatro_descargas(self):
        self.vista("Catálogo")
        self.assertEqual(sum(e.label.startswith("SIM-") for e in self.app.expander), 45)
        self.vista("Casos")
        self.boton("Abrir ejemplo sintético").click().run()
        self.assertFalse(self.app.exception)
        self.vista("Alternativas")
        for campo in self.app.number_input:
            campo.set_value(1.0)
        next(c for c in self.app.checkbox if c.label.startswith("Revisé")).check().run()
        self.boton("Confirmar solución y cantidades").click().run()
        self.assertFalse(self.app.exception)
        self.vista("Cierre")
        self.assertIn("$12,100.00", [m.value for m in self.app.metric])
        self.boton("Preparar PDF y Word").click().run()
        self.assertFalse(self.app.exception)
        self.assertEqual(len(self.app.get("download_button")), 4)
        self.mock.assert_not_called()

    def test_error_proveedor_conserva_texto_y_ficha(self):
        antes = deepcopy(self.app.session_state.ficha)
        self.app.text_area[0].input("Tenemos diez usuarios").run()
        self.boton("Incorporar a la ficha").click().run()
        self.assertFalse(self.app.exception)
        self.assertTrue(self.app.error)
        self.assertEqual(self.app.text_area[0].value, "Tenemos diez usuarios")
        self.assertEqual(self.app.session_state.ficha, antes)

    def test_caso_nuevo_analizado_y_recuperado(self):
        self.vista("Casos")
        next(c for c in self.app.text_input if c.label == "Empresa").input("Empresa nueva")
        self.boton("Crear caso").click().run()
        self.assertEqual(self.app.session_state.ficha["empresa"], "Empresa nueva")
        case_id = self.app.session_state.ficha["case_id"]
        # Nueva sesión: comprueba recuperación de disco. AppTest conserva nodos del formulario
        # anterior cuando este provoca st.rerun después de emitir sus widgets.
        self.app = AppTest.from_file(str(Path(__file__).parents[1] / "app.py"), default_timeout=20)
        self.app.session_state.ficha = F.cargar(case_id)
        self.app.run()
        self.mock.side_effect = None
        self.mock.return_value = ({"hechos": [{"clave": "usuarios_m365", "valor": 10, "cita": "Tenemos diez usuarios"}],
                                   "candidatos": [{"sku": "SIM-M365-02", "cantidad": 10, "hecho_cantidad": "usuarios_m365"}]}, {})
        self.app.text_area[0].input("Tenemos diez usuarios")
        self.boton("Incorporar a la ficha").click().run()
        self.assertFalse(self.app.exception)
        self.assertEqual(self.app.session_state.ficha["revision"], 1)
        self.vista("Casos")
        self.app.selectbox[0].select("CASO-02")
        self.boton("Abrir caso").click().run()
        self.vista("Casos")
        self.app.selectbox[0].select(case_id)
        self.boton("Abrir caso").click().run()
        self.assertEqual(self.app.session_state.ficha["hechos"][0]["valor"], 10)
        self.assertEqual(len(self.app.session_state.ficha["transcript"]), 1)


if __name__ == "__main__":
    unittest.main()
