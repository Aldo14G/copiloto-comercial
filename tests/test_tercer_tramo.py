"""Tercer tramo: del caso confirmado a los entregables, su recuperación y lo que la app afirma.

Todo corre con proveedor simulado: ninguna prueba llama a Claude."""
import json
import re
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader
from streamlit.testing.v1 import AppTest

from copiloto import catalog, claude_engine as C, ficha as F
from copiloto.demo import ejemplo
from copiloto.export import archivos

RAIZ = Path(__file__).parents[1]
CATALOGO = json.loads((RAIZ / "data/catalogo.json").read_text(encoding="utf-8"))
PRECIOS = {p["sku"]: p for p in CATALOGO["productos"]}
# Afirmaciones de impacto que la app no puede sustentar con sus datos.
IMPACTO = re.compile(r"reduc|ahorr|semanas? menos|\d+\s*%\s*(menos|más)|\d+x más", re.I)


def _texto_pdf(datos):
    return "\n".join(p.extract_text() for p in PdfReader(BytesIO(datos)).pages)


def _texto_docx(datos):
    doc = Document(BytesIO(datos))
    celdas = [c.text for t in doc.tables for fila in t.rows for c in fila.cells]
    return "\n".join([p.text for p in doc.paragraphs] + celdas)


class _Disco(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.storage = patch.object(F, "FICHAS", Path(self.tmp.name))
        self.storage.start()
        self.proveedor = patch.object(C, "ask", side_effect=AssertionError("no se permiten llamadas al proveedor"))
        self.mock = self.proveedor.start()

    def tearDown(self):
        self.proveedor.stop()
        self.storage.stop()
        self.tmp.cleanup()


class TestCierreDePuntaAPunta(_Disco):
    def _cerrar(self):
        f = ejemplo()
        F.confirmar(f, [{"sku": "SIM-DC-01", "cantidad": 1}, {"sku": "SIM-DC-08", "cantidad": 1}], 12)
        return f, F.cotizacion(f, usar_confirmacion=True)

    def test_importes_salen_del_catalogo_y_llegan_iguales_a_pdf_y_word(self):
        f, q = self._cerrar()
        mensual = sum(PRECIOS[s]["precio_mensual_mxn"] for s in ("SIM-DC-01", "SIM-DC-08"))
        unico = sum(PRECIOS[s]["cargo_unico_mxn"] for s in ("SIM-DC-01", "SIM-DC-08"))
        self.assertEqual((q["totales"]["mensual"], q["totales"]["unico"]), (mensual, unico))
        self.assertEqual((mensual, unico), (12100, 11000))
        self.assertTrue(q["parcial"])
        self.assertIn("ESC-07", q["escalamientos"])
        salida = archivos(f, q)
        self.assertEqual(sorted(salida), sorted(f"{f['case_id']}-r{f['revision']}-{t}.{e}"
                                                for t in ("propuesta", "expediente") for e in ("pdf", "docx")))
        for nombre, datos in salida.items():
            texto = _texto_pdf(datos) if nombre.endswith(".pdf") else _texto_docx(datos)
            self.assertIn(f["case_id"], texto, nombre)
            self.assertIn("ESC-07", texto, nombre)
            self.assertIn("DATOS SIMULADOS", texto, nombre)
            if "-propuesta." in nombre:
                self.assertIn("$12,100.00", texto, nombre)
                self.assertIn("$11,000.00", texto, nombre)
                self.assertIn("Subtotal parcial", texto, nombre)
        self.mock.assert_not_called()

    def test_caso_recuperado_de_disco_conserva_confirmacion_importe_y_entregables(self):
        f, q = self._cerrar()
        recuperada = F.cargar(f["case_id"])
        self.assertEqual(F.estado_confirmacion(recuperada), "vigente")
        q2 = F.cotizacion(recuperada, usar_confirmacion=True)
        self.assertEqual((q2["totales"], q2["huella"], q2["revision"]), (q["totales"], q["huella"], q["revision"]))
        self.assertEqual(recuperada["transcript"], f["transcript"])
        self.assertEqual(len(archivos(recuperada, q2)), 4)

    def test_cambio_posterior_exige_nueva_confirmacion_antes_de_la_propuesta(self):
        f, _ = self._cerrar()
        F.cambiar_plazo(f, 24)
        q = F.cotizacion(f, usar_confirmacion=True)
        self.assertEqual(q["confirmacion"], "desactualizada")
        self.assertFalse(q["puede_generar_borrador"])
        self.assertEqual([n.rsplit("-", 1)[1] for n in archivos(f, q)], ["expediente.pdf", "expediente.docx"])


class TestProveedorNoInventaPrecios(_Disco):
    def test_precios_y_skus_del_proveedor_se_ignoran(self):
        respuesta = {
            "hechos": [{"clave": "racks", "valor": 1, "cita": "necesitamos un rack"}],
            "candidatos": [
                {"sku": "SIM-DC-01", "cantidad": 1, "hecho_cantidad": "racks",
                 "precio_mensual_mxn": 1, "mensual": 1, "total": 1, "descuento": 0.9},
                {"sku": "SIM-XX-99", "cantidad": 1, "precio_mensual_mxn": 10}],
            "alternativa": {"skus": ["SIM-ZZ-01"], "resumen": "Oferta especial a $1"}}
        f = F.nueva(empresa="Proveedor simulado")
        self.mock.side_effect = None
        self.mock.return_value = (respuesta, {"segundos": 0})
        F.procesar_mensaje(f, "Hoy necesitamos un rack para el servidor.")
        self.assertEqual([c["sku"] for c in f["candidatos"]], ["SIM-DC-01"])
        self.assertIsNone(f["alternativa"])
        self.assertTrue(any("SIM-XX-99" in a for a in f["alertas_validacion"]))
        F.confirmar(f, [{"sku": c["sku"], "cantidad": c["cantidad"]} for c in f["candidatos"]])
        q = F.cotizacion(f, usar_confirmacion=True)
        self.assertEqual([l["sku"] for l in q["lineas"]], ["SIM-DC-01"])
        self.assertEqual(q["totales"]["mensual"], PRECIOS["SIM-DC-01"]["precio_mensual_mxn"])
        self.assertEqual(q["totales"]["unico"], PRECIOS["SIM-DC-01"]["cargo_unico_mxn"])

    def test_cantidad_sin_cita_numerica_no_llega_al_importe(self):
        f = F.nueva(empresa="Proveedor simulado")
        self.mock.side_effect = None
        self.mock.return_value = ({"hechos": [], "candidatos": [{"sku": "SIM-DC-01", "cantidad": 4}]}, {})
        F.procesar_mensaje(f, "Queremos sacar el servidor de la oficina.")
        self.assertIsNone(f["candidatos"][0]["cantidad"])
        q = F.cotizacion(f)
        self.assertFalse(q["puede_generar_borrador"])


class TestLoQueMuestraLaApp(_Disco):
    def setUp(self):
        super().setUp()
        self.app = AppTest.from_file(str(RAIZ / "app.py"), default_timeout=20).run()

    def vista(self, nombre):
        self.app.radio(key="navegacion").set_value(nombre).run()
        self.assertFalse(self.app.exception)

    def textos(self):
        grupos = ("title", "header", "subheader", "markdown", "caption", "info", "warning", "error", "success")
        return [str(e.value) for g in grupos for e in getattr(self.app, g)]

    def test_conteos_del_catalogo_se_calculan_y_no_hay_afirmaciones_de_impacto(self):
        esperado = f"{len(CATALOGO['productos'])} productos"
        self.assertTrue(any(esperado in t for t in self.textos()))
        self.vista("Casos")
        next(b for b in self.app.button if b.label == "Abrir ejemplo sintético").click().run()
        # Catálogo cita las fichas de producto tal cual; se revisa el texto propio de las vistas de trabajo.
        for nombre in ("Casos", "Reunión", "Alternativas", "Cierre", "Historial"):
            self.vista(nombre)
            for texto in self.textos():
                self.assertIsNone(IMPACTO.search(texto), f"{nombre}: {texto}")
        self.vista("Catálogo")
        self.assertTrue(any(f"{len(CATALOGO['familias'])} familias · {len(CATALOGO['paquetes'])} paquetes" in t
                            for t in self.textos()))
        self.mock.assert_not_called()

    def test_cierre_muestra_importes_del_cotizador_y_estado_legible(self):
        self.vista("Casos")
        next(b for b in self.app.button if b.label == "Abrir ejemplo sintético").click().run()
        self.vista("Cierre")
        self.assertEqual([m.value for m in self.app.metric], [])
        self.assertTrue(any("confirmación sin confirmar" in t for t in self.textos()))
        self.vista("Alternativas")
        for campo in self.app.number_input:
            campo.set_value(1.0)
        next(c for c in self.app.checkbox if c.label.startswith("Revisé")).check()
        next(b for b in self.app.button if b.label == "Confirmar solución y cantidades").click().run()
        self.vista("Cierre")
        q = F.cotizacion(self.app.session_state.ficha, usar_confirmacion=True)
        self.assertEqual([m.value for m in self.app.metric],
                         [f"${q['totales']['mensual']:,.2f}", f"${q['totales']['unico']:,.2f}"])
        next(b for b in self.app.button if b.label == "Preparar PDF y Word").click().run()
        self.assertEqual(sorted(b.label for b in self.app.get("download_button")),
                         ["Descargar expediente · PDF", "Descargar expediente · Word",
                          "Descargar propuesta · PDF", "Descargar propuesta · Word"])
        self.vista("Historial")
        self.assertTrue(any("Queremos conservar los dos servidores" in t for t in self.textos()))
        self.mock.assert_not_called()


if __name__ == "__main__":
    unittest.main()
