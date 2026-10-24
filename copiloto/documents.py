"""Un contenido editable, con representaciones Markdown, Word y PDF."""
from dataclasses import dataclass
from functools import lru_cache
from html import escape
from io import BytesIO
from pathlib import Path
import re


@dataclass(frozen=True)
class Texto:
    texto: str
    nivel: int = 0


@dataclass(frozen=True)
class Tabla:
    columnas: tuple[str, ...]
    filas: list[tuple[str, ...]]


@dataclass(frozen=True)
class Documento:
    titulo: str
    identificacion: str
    contenido: list[Texto | Tabla]


def _texto(valor):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(valor))


def markdown(doc):
    out = [f"# {doc.titulo}", doc.identificacion]
    for bloque in doc.contenido:
        if isinstance(bloque, Texto):
            out.append(("## " if bloque.nivel else "") + _texto(bloque.texto).replace("$", "\\$"))
        else:
            filas = [bloque.columnas, tuple("---" for _ in bloque.columnas), *bloque.filas]
            out.extend("| " + " | ".join(_texto(c).replace("|", "\\|").replace("$", "\\$").replace("\n", "<br>")
                                         for c in fila) + " |" for fila in filas)
    return "\n\n".join(out)


def word(doc):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    archivo = Document()
    seccion = archivo.sections[0]
    seccion.page_width, seccion.page_height = Cm(21), Cm(29.7)
    seccion.top_margin = seccion.bottom_margin = Cm(1.9)
    seccion.left_margin = seccion.right_margin = Cm(1.8)
    for nombre, tamano in (("Normal", 10), ("Title", 24), ("Heading 1", 14), ("Caption", 8)):
        estilo = archivo.styles[nombre]
        fuente = estilo.font
        fuente.name, fuente.size = "Georgia", Pt(tamano)
        fuente.bold = False
        fuente.color.rgb = RGBColor.from_string("000000" if nombre == "Title" else "19382D")
        rfonts = estilo.element.rPr.rFonts
        for atributo in list(rfonts.attrib):
            if "theme" in atributo.lower():
                del rfonts.attrib[atributo]
        if estilo.element.pPr is not None:
            for borde in list(estilo.element.pPr.findall(qn("w:pBdr"))):
                estilo.element.pPr.remove(borde)
        estilo.paragraph_format.space_after = Pt(6)
        if nombre == "Heading 1":
            estilo.paragraph_format.space_before = Pt(12)
    seccion.header.paragraphs[0].text = "COPILOTO COMERCIAL / DATOS SIMULADOS"
    archivo.add_paragraph(_texto(doc.titulo), "Title")
    archivo.add_paragraph(_texto(doc.identificacion).replace("`", ""), "Caption")
    for bloque in doc.contenido:
        if isinstance(bloque, Texto):
            archivo.add_paragraph(_texto(bloque.texto), "Heading 1" if bloque.nivel else "Normal")
            continue
        tabla = archivo.add_table(rows=1, cols=len(bloque.columnas))
        tabla.style = "Table Grid"
        margenes = OxmlElement("w:tblCellMar")
        for lado in ("top", "left", "bottom", "right"):
            margen = OxmlElement(f"w:{lado}")
            margen.set(qn("w:w"), "80")
            margen.set(qn("w:type"), "dxa")
            margenes.append(margen)
        tabla._tbl.tblPr.append(margenes)
        for celda, valor in zip(tabla.rows[0].cells, bloque.columnas):
            celda.text = _texto(valor)
            sombreado = OxmlElement("w:shd")
            sombreado.set(qn("w:fill"), "E7EEE8")
            celda._tc.get_or_add_tcPr().append(sombreado)
        tabla.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for fila in bloque.filas:
            for celda, valor in zip(tabla.add_row().cells, fila):
                celda.text = _texto(valor)
        for fila in tabla.rows:
            for celda in fila.cells:
                for parrafo in celda.paragraphs:
                    for run in parrafo.runs:
                        run.font.size = Pt(9)
    pie = seccion.footer.paragraphs[0]
    pie.text = "Borrador para revisión / Página "
    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    pie._p.append(campo)
    salida = BytesIO()
    archivo.save(salida)
    return salida.getvalue()


@lru_cache(maxsize=1)
def _fuente_pdf():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    ruta = Path(__file__).parents[1] / "static/fonts/SourceSerif4.ttf"
    pdfmetrics.registerFont(TTFont("DocumentoSerif", str(ruta)))
    return "DocumentoSerif"


def pdf(doc):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle

    fuente = _fuente_pdf()
    estilos = {nombre: ParagraphStyle(nombre, fontName=fuente, fontSize=tamano, leading=interlinea,
                                     spaceAfter=8, textColor=colors.HexColor("#19382D"))
               for nombre, tamano, interlinea in (("cuerpo", 10, 14), ("titulo", 23, 28),
                                                 ("seccion", 14, 18), ("pequeno", 8, 11), ("celda", 9, 12))}
    def parrafo(texto, estilo="cuerpo"):
        return Paragraph(escape(_texto(texto).replace("`", "")).replace("\n", "<br/>"), estilos[estilo])
    flujo = [parrafo(doc.titulo, "titulo"), parrafo(doc.identificacion, "pequeno")]
    for bloque in doc.contenido:
        if isinstance(bloque, Texto):
            flujo.append(parrafo(bloque.texto, "seccion" if bloque.nivel else "cuerpo"))
            continue
        tabla = LongTable([[parrafo(c, "celda") for c in fila] for fila in [bloque.columnas, *bloque.filas]],
                          colWidths=[17.4 * cm / len(bloque.columnas)] * len(bloque.columnas),
                          repeatRows=1, splitInRow=1)
        tabla.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7EEE8")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("LINEBELOW", (0, 0), (-1, -1), .3, colors.HexColor("#CED8D0")),
                                   ("TOPPADDING", (0, 0), (-1, -1), 6),
                                   ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        flujo.extend([tabla, Spacer(1, 10)])
    def pagina(canvas, documento):
        canvas.setFont(fuente, 8)
        canvas.drawString(1.8 * cm, 28.5 * cm, "COPILOTO COMERCIAL / DATOS SIMULADOS")
        canvas.drawString(1.8 * cm, 1 * cm, f"Borrador para revisión / Página {documento.page}")
    salida = BytesIO()
    SimpleDocTemplate(salida, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm,
                      topMargin=1.9 * cm, bottomMargin=1.9 * cm).build(flujo, onFirstPage=pagina, onLaterPages=pagina)
    return salida.getvalue()
